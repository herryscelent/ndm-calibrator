"""Excel export: one .xlsx workbook with summary, parameters, metrics, etc."""

from __future__ import annotations

from pathlib import Path
from typing import Dict

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from core.exceptions import NDMCError
from core.project import Project


def _write_sheet(wb, title: str, headers: list, rows: list, bold_first_col=False):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    # Auto-width (heuristic).
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[letter].width = min(max_len + 4, 40)
    return ws


def export_excel(project: Project, dest_dir: Path | str, checks: Dict[str, bool]) -> str:
    dest_dir = Path(dest_dir)
    if not dest_dir.exists():
        raise NDMCError(f"目标目录不存在: {dest_dir}")
    wb = Workbook()
    wb.remove(wb.active)

    results = project.results

    if checks.get("summary"):
        rows = [
            ("项目名称", project.name),
            ("描述", project.description),
            ("创建时间", project.created_at),
            ("修改时间", project.modified_at),
            ("数据来源", project.data_source or ""),
            ("模型状态", ",".join(project.model.states) if project.model else ""),
            ("参数", ",".join(project.parameters.names)),
        ]
        _write_sheet(wb, "Project Summary", ["项", "值"], rows)

    if checks.get("parameters"):
        rows = [[p.name, p.initial, p.lower, p.upper, "待辨识" if p.free else "固定"] for p in project.parameters]
        _write_sheet(wb, "Parameters", ["参数", "初值", "下限", "上限", "状态"], rows)

    if checks.get("metrics"):
        cal = results.get("calibration", {})
        metrics = cal.get("metrics", {})
        rows = []
        for output in metrics:
            for k, v in metrics[output].items():
                rows.append([output, k, v if isinstance(v, (int, float)) else str(v)])
        _write_sheet(wb, "Metrics", ["输出", "指标", "值"], rows)

    if checks.get("simulation"):
        ctx = _make_prediction(project)
        if ctx is not None:
            rows = [["t", *[f"{out}_sim" for out in project.model.outputs]]]
            for i, t in enumerate(project.experiment.time):
                rows.append([t, *[ctx["predicted"][o][i] for o in project.model.outputs]])
            _write_sheet(wb, "Simulation", rows[0], rows[1:])

    if checks.get("residuals"):
        cal = results.get("calibration", {})
        for out in project.model.outputs if project.model else []:
            res = cal.get("residuals", {}).get(out, [])
            stats = cal.get("residual_stats", {}).get(out, {})
            rows = [[k, v] for k, v in stats.items()] + [["residual", r] for r in res]
            _write_sheet(wb, f"Residuals_{out}", ["项", "值"], rows)

    if checks.get("diagnostics"):
        cal = results.get("calibration", {})
        ident = cal.get("identifiability", {})
        rows = [
            ("condition_number", ident.get("condition_number")),
            ("rank", ident.get("rank")),
            ("n_samples", ident.get("n_samples")),
            ("n_parameters", ident.get("n_parameters")),
        ]
        for i, sv in enumerate(ident.get("singular_values", [])):
            rows.append((f"singular_value_{i+1}", sv))
        for w in ident.get("warnings", []):
            rows.append(("warning", w))
        _write_sheet(wb, "Identifiability", ["项", "值"], rows)

    if checks.get("bootstrap"):
        boot = results.get("bootstrap", {})
        rows = []
        for name, info in boot.get("parameters", {}).items():
            rows.append([
                name, info.get("mean"), info.get("median"),
                info.get("p2_5"), info.get("p97_5"), info.get("std"),
            ])
        _write_sheet(wb, "Bootstrap", ["参数", "均值", "中位数", "2.5%", "97.5%", "标准差"], rows)

    out_path = dest_dir / f"{project.name}_results.xlsx"
    wb.save(out_path)
    return str(out_path)


def _make_prediction(project):
    """Return predicted outputs for the current optimal parameters."""
    if project.model is None or project.experiment is None:
        return None
    entry = project.results.get("identification", {})
    params = entry.get("result", {}).get("parameters")
    if not params:
        return None
    try:
        from identification.base import build_context

        ctx = build_context(project)
        free_names = ctx.free_names
        theta = np.array([params[n] for n in free_names], dtype=float)
        predicted = ctx.predict_outputs(theta)
        return {"predicted": predicted}
    except Exception:
        return None
